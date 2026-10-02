"""Classificação de documentos por emissor com o JEV (TypeSafe AI, modelo System One).

O JEV não gera texto: recebe um estado e perguntas tipadas e devolve escolhas,
notas e probabilidades com confiança. Aqui ele lê o TÍTULO de cada documento
(fato relevante, comunicado, aviso, notícia, análise) e responde três perguntas:

  evento     (choice) : tipo do evento
  impacto    (score)  : negativo / neutro / positivo para o credor
  relevante  (noul)   : o documento trata do crédito deste emissor?

Modos:
  python analises/jev_classificar.py amostra   # sorteia 50 docs, gera a planilha de validação (sem respostas do JEV)
  python analises/jev_classificar.py testar    # verifica a chave (GET /v1/models) e faz 1 chamada
  python analises/jev_classificar.py validar   # classifica os 50 da amostra (respostas guardadas à parte)
  python analises/jev_classificar.py comparar  # compara com a planilha preenchida pelo Alisson
  python analises/jev_classificar.py tudo      # classifica todos os documentos (só depois da validação)

A chave vem da variável de ambiente TYPESAFE_API_KEY (segredo do GitHub). Nunca é impressa.
"""

import csv
import json
import os
import random
import re
import sys
import time
import urllib.error
import urllib.request
from pathlib import Path

RAIZ = Path(__file__).resolve().parent.parent
DOCS = RAIZ / "dados" / "derivados" / "documentos.json"
UNIVERSO = RAIZ / "dados" / "referencia" / "universo_series.csv"
PASTA = RAIZ / "dados" / "derivados" / "jev"
AMOSTRA = PASTA / "amostra_validacao.json"
PLANILHA = RAIZ / "validacao" / "jev_validacao_50.xlsx"
API = "https://api.typesafe.ai/v1"
MODELO = "jev-latest"

EVENTOS = {
    "credito_negativo": "Evento de crédito negativo: inadimplemento, vencimento antecipado, recuperação judicial, waiver, rebaixamento de rating, renegociação de dívida",
    "nova_divida": "Nova emissão ou captação de dívida (debêntures, empréstimos, financiamentos)",
    "pagamento_divida": "Pagamento de juros, amortização ou resgate de dívida já existente",
    "aquisicao_venda": "Aquisição, venda de ativos ou participações, leilão vencido, fusão",
    "operacional": "Operação dos ativos: energização, entrada em operação, licença, obra, reajuste de receita (RAP)",
    "societario_governanca": "Mudança de controle, reorganização societária, eleição de diretores ou conselheiros",
    "resultado": "Resultados financeiros, guidance, proventos, dividendos",
    "rating": "Atribuição ou manutenção de rating de crédito",
    "analise_credito": "Análise ou relatório de crédito sobre a debênture ou a dívida do emissor, feito por casa de análise",
    "outro": "Outro assunto, sem relação clara com crédito",
}
IMPACTO = ["negativo para o credor", "neutro para o credor", "positivo para o credor"]
PERGUNTAS = {
    "evento": {"type": "choice", "instructions": "Qual é o tipo de evento descrito neste documento?", "criteria": EVENTOS},
    "impacto": {"type": "score", "instructions": "Para quem detém debêntures deste emissor, qual o impacto do evento sobre a capacidade de pagamento?", "criteria": IMPACTO},
    "relevante": {"type": "noul", "instructions": (
        "Este documento é principalmente sobre este emissor e afeta, mesmo que indiretamente, sua capacidade de pagar dívidas? "
        "Conta como sim: dívida, pagamentos, resultados, receita, entrada em operação de ativos, leilões, aquisições, venda de ativos, "
        "mudança de controle. Conta como não: o emissor é só uma entre várias empresas citadas, ou o assunto é apenas governança "
        "interna (eleição, comitês) sem efeito financeiro.")},
}
# títulos sem conteúdo: o JEV respondeu com confiança alta e errado na validação (02/10/2026); não são enviados
GENERICOS = re.compile(r"^\s*(comunicado( ao mercado)?|fato relevante|aviso aos (acionistas|debenturistas)|outros comunicados|"
                       r"esclarecimentos?( sobre .{0,20})?)\s*[.:-]?\s*$", re.I)
CONFIANCA_MINIMA = 0.8  # validação: 89% de acerto no tipo de evento com confiança >= 0,8, 36% abaixo disso

TIPOS_CLASSIFICADOS = ("fato_relevante", "comunicado", "aviso_debenturistas", "noticia", "analise")
NOME_TIPO = {"fato_relevante": "Fato relevante (CVM)", "comunicado": "Comunicado ao mercado (CVM)",
             "aviso_debenturistas": "Aviso aos debenturistas (CVM)", "noticia": "Notícia", "analise": "Análise de casa de research"}


def chamar(caminho: str, corpo: dict | None = None) -> dict:
    chave = os.environ.get("TYPESAFE_API_KEY")
    if not chave:
        raise SystemExit("TYPESAFE_API_KEY não definida")
    req = urllib.request.Request(
        f"{API}/{caminho}", method="POST" if corpo is not None else "GET",
        data=json.dumps(corpo).encode() if corpo is not None else None,
        headers={"Authorization": f"Bearer {chave}", "Content-Type": "application/json"},
    )
    for tentativa in range(4):
        try:
            with urllib.request.urlopen(req, timeout=60) as r:
                return json.load(r)
        except urllib.error.HTTPError as e:
            if e.code in (429, 500, 502, 503) and tentativa < 3:
                time.sleep(2 ** tentativa * 2)
                continue
            raise SystemExit(f"API respondeu {e.code}: {e.read()[:300].decode(errors='ignore')}")


def estado(doc: dict, emissor: str) -> str:
    return (f"Emissor de debêntures: {emissor}. Setor: transmissão de energia elétrica no Brasil. "
            f"Documento: {NOME_TIPO[doc['tipo']]}. Fonte: {doc.get('fonte', '')}. Data: {doc.get('data', '')}. "
            f"Título: {doc['titulo']}")


def classificar(doc: dict, emissor: str) -> dict:
    r = chamar("systemone", {"model": MODELO, "state": estado(doc, emissor), "questions": PERGUNTAS})
    a = r["answers"]
    return {
        "evento": a["evento"]["choice"], "evento_confianca": a["evento"]["confidence"],
        "evento_probabilidades": a["evento"]["probabilities"],
        "impacto_score": a["impacto"]["score"], "impacto_confianca": a["impacto"]["confidence"],
        "relevante": a["relevante"]["noul"], "modelo": r.get("model"), "tokens": r.get("usage", {}).get("input_tokens"),
    }


def todos_docs() -> list[dict]:
    emissores = {u["cnpj"]: u["emissor_atual_snd"] for u in csv.DictReader(UNIVERSO.open(encoding="utf-8"))}
    base = json.loads(DOCS.read_text(encoding="utf-8"))["emissores"]
    saida = []
    for cnpj, ls in base.items():
        for d in ls:
            if d["tipo"] in TIPOS_CLASSIFICADOS:
                saida.append({**d, "cnpj": cnpj, "emissor": emissores.get(cnpj, cnpj).title(), "id": f"{cnpj}|{d['tipo']}|{d.get('data','')}|{d['titulo'][:60]}"})
    return saida


def gerar_amostra() -> None:
    from openpyxl import Workbook
    from openpyxl.styles import Alignment, Font, PatternFill
    from openpyxl.worksheet.datavalidation import DataValidation

    docs = todos_docs()
    random.seed(20261002)
    alvo = {"fato_relevante": 14, "comunicado": 10, "aviso_debenturistas": 5, "noticia": 13, "analise": 8}
    amostra = []
    for tipo, n in alvo.items():
        grupo = [d for d in docs if d["tipo"] == tipo]
        amostra += random.sample(grupo, min(n, len(grupo)))
    PASTA.mkdir(parents=True, exist_ok=True)
    AMOSTRA.write_text(json.dumps(amostra, ensure_ascii=False, indent=1), encoding="utf-8")

    wb = Workbook()
    ws = wb.active
    ws.title = "Classificar"
    cab = ["nº", "Emissor", "Tipo de documento", "Data", "Fonte", "Título", "Evento", "Impacto para o credor", "Trata do crédito do emissor?", "Comentário"]
    ws.append(cab)
    for i, d in enumerate(amostra, 1):
        ws.append([i, d["emissor"], NOME_TIPO[d["tipo"]], d.get("data", ""), d.get("fonte", ""), d["titulo"], "", "", "", ""])
        ws.cell(row=i + 1, column=6).hyperlink = d.get("url") or None
    for c in ws[1]:
        c.font = Font(bold=True, color="FFFFFF")
        c.fill = PatternFill("solid", fgColor="1C5CAB")
        c.alignment = Alignment(vertical="center", wrap_text=True)
    larguras = [5, 32, 26, 11, 16, 70, 24, 24, 16, 30]
    for i, w in enumerate(larguras):
        ws.column_dimensions[chr(65 + i)].width = w
    for linha in ws.iter_rows(min_row=2):
        for c in linha:
            c.alignment = Alignment(vertical="top", wrap_text=True)
        for col in (7, 8, 9):
            linha[col - 1].fill = PatternFill("solid", fgColor="FFF7D6")
    n = len(amostra) + 1
    dv_ev = DataValidation(type="list", formula1='"' + ",".join(EVENTOS) + '"', allow_blank=True)
    dv_im = DataValidation(type="list", formula1='"negativo,neutro,positivo"', allow_blank=True)
    dv_re = DataValidation(type="list", formula1='"sim,não"', allow_blank=True)
    for dv, col in ((dv_ev, "G"), (dv_im, "H"), (dv_re, "I")):
        ws.add_data_validation(dv)
        dv.add(f"{col}2:{col}{n}")
    ws.freeze_panes = "B2"

    leg = wb.create_sheet("Legenda")
    leg.append(["Evento", "O que significa"])
    for k, v in EVENTOS.items():
        leg.append([k, v])
    leg.append([])
    leg.append(["Impacto", "negativo / neutro / positivo para quem detém debêntures do emissor"])
    leg.append(["Trata do crédito?", "sim = o documento é sobre a dívida ou a situação financeira deste emissor; não = só cita o nome"])
    leg.column_dimensions["A"].width = 22
    leg.column_dimensions["B"].width = 110
    for c in leg[1]:
        c.font = Font(bold=True)
    PLANILHA.parent.mkdir(parents=True, exist_ok=True)
    wb.save(PLANILHA)
    print(f"{len(amostra)} documentos sorteados -> {PLANILHA.relative_to(RAIZ)}")


def testar() -> None:
    modelos = chamar("models")
    print("chave válida; modelos:", [m.get("name") for m in modelos.get("models", [])])
    d = todos_docs()[0]
    r = classificar(d, d["emissor"])
    print("chamada de teste ok:", d["titulo"][:70], "->", r["evento"], round(r["evento_confianca"], 2), "| impacto", round(r["impacto_score"], 2))


def validar() -> None:
    amostra = json.loads(AMOSTRA.read_text(encoding="utf-8"))
    respostas = {}
    for d in amostra:
        respostas[d["id"]] = classificar(d, d["emissor"])
        time.sleep(0.3)
    (PASTA / "respostas_amostra.json").write_text(json.dumps(respostas, ensure_ascii=False, indent=1), encoding="utf-8")
    print(f"{len(respostas)} documentos classificados pelo JEV (respostas guardadas, ainda não comparadas)")


def comparar() -> None:
    from openpyxl import load_workbook
    amostra = json.loads(AMOSTRA.read_text(encoding="utf-8"))
    jev = json.loads((PASTA / "respostas_amostra.json").read_text(encoding="utf-8"))
    ws = load_workbook(PLANILHA)["Classificar"]
    humano = {}
    for linha in ws.iter_rows(min_row=2, values_only=True):
        if linha[0]:
            humano[int(linha[0])] = {"evento": linha[6], "impacto": linha[7], "relevante": linha[8]}
    nivel = lambda s: 0 if s < 0.67 else 1 if s < 1.33 else 2
    mapa_imp = {"negativo": 0, "neutro": 1, "positivo": 2}
    linhas, acertos = [], {"evento": [], "impacto": [], "relevante": []}
    for i, d in enumerate(amostra, 1):
        h, j = humano.get(i, {}), jev.get(d["id"])
        if not j or not h.get("evento"):
            continue
        ok_ev = h["evento"] == j["evento"]
        ok_im = mapa_imp.get(str(h["impacto"]).strip().lower()) == nivel(j["impacto_score"]) if h["impacto"] else None
        ok_re = (str(h["relevante"]).strip().lower() == "sim") == (j["relevante"] >= .5) if h["relevante"] else None
        acertos["evento"].append((ok_ev, j["evento_confianca"]))
        if ok_im is not None:
            acertos["impacto"].append((ok_im, j["impacto_confianca"]))
        if ok_re is not None:
            acertos["relevante"].append((ok_re, max(j["relevante"], 1 - j["relevante"])))
        linhas.append({"n": i, "titulo": d["titulo"][:80], "humano_evento": h["evento"], "jev_evento": j["evento"],
                       "conf": round(j["evento_confianca"], 2), "ok": ok_ev})
    print(f"documentos comparados: {len(linhas)}")
    for k, v in acertos.items():
        if not v:
            continue
        taxa = sum(o for o, _ in v) / len(v)
        alta = [o for o, c in v if c >= .8]
        baixa = [o for o, c in v if c < .8]
        print(f"{k:10} acerto {taxa:.0%} (n={len(v)}) | confiança >= 0,8: {sum(alta)}/{len(alta)} | < 0,8: {sum(baixa)}/{len(baixa)}")
    (PASTA / "comparacao.json").write_text(json.dumps({"linhas": linhas}, ensure_ascii=False, indent=1), encoding="utf-8")


def nivel_impacto(score: float) -> str:
    return "negativo" if score < 0.67 else "neutro" if score < 1.33 else "positivo"


def tudo() -> None:
    """Classifica documentos novos. Confiança >= CONFIANCA_MINIMA entra direto; abaixo vai para a fila de revisão."""
    saida = {}
    arq = PASTA / "classificacao.json"
    if arq.exists():
        saida = json.loads(arq.read_text(encoding="utf-8"))
    novos = genericos = 0
    for d in todos_docs():
        if d["id"] in saida:
            continue
        if GENERICOS.match(d["titulo"]):
            saida[d["id"]] = {"status": "titulo_generico"}
            genericos += 1
            continue
        r = classificar(d, d["emissor"])
        r["impacto"] = nivel_impacto(r["impacto_score"])
        r["status"] = "automatico" if r["evento_confianca"] >= CONFIANCA_MINIMA else "a_revisar"
        saida[d["id"]] = r
        novos += 1
        time.sleep(0.3)
    arq.write_text(json.dumps(saida, ensure_ascii=False, indent=1), encoding="utf-8")

    # fila de calibração: o que o JEV não teve certeza, para o Alisson decidir aos poucos
    docs = {d["id"]: d for d in todos_docs()}
    fila = [(i, r) for i, r in saida.items() if r.get("status") == "a_revisar" and i in docs]
    with (PASTA / "fila_revisao.csv").open("w", newline="", encoding="utf-8") as f:
        w = csv.writer(f, lineterminator="\n")
        w.writerow(["id", "emissor", "tipo", "data", "titulo", "jev_evento", "jev_confianca", "jev_impacto", "decisao_evento", "decisao_impacto"])
        for i, r in sorted(fila, key=lambda x: x[1]["evento_confianca"]):
            d = docs[i]
            w.writerow([i, d["emissor"], d["tipo"], d.get("data", ""), d["titulo"], r["evento"], round(r["evento_confianca"], 2), r["impacto"], "", ""])
    from collections import Counter
    st = Counter(r.get("status") for r in saida.values())
    print(f"{novos} novos classificados, {genericos} títulos genéricos ignorados; total {dict(st)}; fila de revisão: {len(fila)}")


if __name__ == "__main__":
    {"amostra": gerar_amostra, "testar": testar, "validar": validar, "comparar": comparar, "tudo": tudo}[sys.argv[1]]()
